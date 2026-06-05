from io import BytesIO

from docx import Document


class DocxParseError(ValueError):
    """Raised when a Word document cannot be parsed into text."""


def parse_docx(file_bytes: bytes) -> str:
    """Extract paragraph and table text from a .docx file."""
    if not file_bytes:
        raise DocxParseError("上传的文件为空。")

    try:
        document = Document(BytesIO(file_bytes))
    except Exception as exc:  # python-docx raises different low-level exceptions for invalid archives.
        raise DocxParseError("无法解析 Word 文件，请确认上传的是有效的 .docx 文件。") from exc

    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append("\t".join(cells))

    text = clean_text("\n".join(blocks))
    if not text:
        raise DocxParseError("未能从 Word 文件中提取到有效文本。")
    return text


def clean_text(text: str) -> str:
    """Normalize script text while preserving line-based screenplay structure."""
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.strip() for line in normalized.split("\n")]

    cleaned: list[str] = []
    previous_blank = False
    for line in lines:
        if not line:
            if not previous_blank and cleaned:
                cleaned.append("")
            previous_blank = True
            continue
        cleaned.append(line)
        previous_blank = False

    return "\n".join(cleaned).strip()
