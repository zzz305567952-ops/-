from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.schemas import StoryboardShot

HEADERS = [
    "镜头序列",
    "景别",
    "拍摄角度",
    "运镜",
    "时长",
    "画面内容",
    "台词",
    "光影氛围",
    "AI绘图提示词",
]


def build_storyboard_workbook(shots: list[StoryboardShot]) -> bytes:
    """Build an Excel workbook from storyboard shots and return .xlsx bytes."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "分镜表"
    worksheet.append(HEADERS)

    for shot in shots:
        worksheet.append(
            [
                shot.sequence,
                shot.shot_size,
                shot.camera_angle,
                shot.camera_movement,
                shot.duration,
                shot.visual_content,
                shot.dialogue,
                shot.lighting_mood,
                shot.ai_image_prompt,
            ]
        )

    _style_worksheet(worksheet)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _style_worksheet(worksheet) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = [12, 12, 14, 14, 10, 36, 32, 36, 56]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
