import pytest

pytest.importorskip("docx", reason="python-docx 未安装；请先运行 pip install -r requirements.txt 后执行 API 测试。")
pytest.importorskip("fastapi", reason="fastapi 未安装；请先运行 pip install -r requirements.txt 后执行 API 测试。")
pytest.importorskip("pydantic", reason="pydantic 未安装；请先运行 pip install -r requirements.txt 后执行 API 测试。")
pytest.importorskip("openpyxl", reason="openpyxl 未安装；请先运行 pip install -r requirements.txt 后执行 Excel 导出测试。")

from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)


def _sample_docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("第1场 王府 长廊 内 夜")
    document.add_paragraph("苏清月站在廊下，雨声淅沥。")
    document.add_paragraph("苏清月：你终究还是来了。")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_create_storyboard_from_docx() -> None:
    response = client.post(
        "/api/storyboards/from-docx",
        files={"file": ("script.docx", _sample_docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )

    assert response.status_code == 200
    data = response.json()
    assert data["filename"] == "script.docx"
    assert len(data["scenes"]) == 1
    assert data["shots"][0]["sequence"] == "1-1"


def test_export_storyboard_excel() -> None:
    payload = {
        "shots": [
            {
                "sequence": "1-1",
                "shot_size": "全景",
                "camera_angle": "平视",
                "camera_movement": "缓慢推进",
                "duration": "3秒",
                "visual_content": "王府长廊雨夜。",
                "dialogue": "",
                "lighting_mood": "冷蓝雨夜。",
                "ai_image_prompt": "古装人物，王府长廊。",
            }
        ]
    }

    response = client.post("/api/storyboards/export-excel", json=payload)

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert response.content.startswith(b"PK")
